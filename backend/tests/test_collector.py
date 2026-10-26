import asyncio
import csv
import io
import time
from datetime import datetime, timezone, timedelta
from email.utils import format_datetime
from unittest.mock import patch
import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func
from app.db import Session, Job, JobItem, Product, Variant
from app.providers import affiliate_product, enrich_ds, MockProvider, OfficialProvider, ProviderError, retry_delay
from app.worker import recover, process_one

def start(client, **kwargs):
    result = client.post('/api/jobs', json={'keyword':'나사', 'end_page':1, 'source':'mock', **kwargs})
    assert result.status_code == 200, result.text
    return result.json()['id']

def wait(client, id, timeout=60):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        job = next(j for j in client.get('/api/jobs').json() if j['id'] == id)
        if job['status'] in ('completed', 'failed', 'completed_with_errors', 'cancelled'):
            return job
        time.sleep(.05)
    raise AssertionError('Job timeout')

def test_one_page_then_ten_pages_and_original_options(client):
    first = wait(client, start(client))
    assert first['status'] == 'completed' and first['counts']['done'] == 50
    rows = client.get('/api/products').json()
    assert rows['total'] == 50 and len(rows['items']) == 20
    detail = client.get(f'/api/products/{rows["items"][0]["id"]}').json()
    assert [v['combination'] for v in detail['variants']] == ['Color:Black;Size:M3','Color:Black;Size:M4','Color:White;Size:M5']
    assert detail['variants'][1]['stock'] == 0
    assert detail['data']['brand'] is None
    assert detail['image_path']
    second = wait(client, start(client, end_page=10))
    assert second['status'] == 'completed'
    assert second['counts']['done'] == 450 and second['counts']['skipped'] == 50
    assert client.get('/api/products').json()['total'] == 500
    assert client.get('/api/duplicates').json()['duplicate_groups'] == []
    with Session() as db:
        assert db.scalar(select(func.count()).select_from(Variant)) == 1500

def test_exports_search_delete_recollect(client):
    wait(client, start(client, page_size=2))
    p = client.get('/api/products?q=스테인리스').json()['items'][0]
    id = p['id']
    csv_response = client.get(f'/api/export/csv?ids={id}')
    rows = list(csv.reader(io.StringIO(csv_response.content.decode('utf-8-sig'))))
    assert len(rows) == 4
    assert 'Black' in rows[1][8]
    xlsx = client.get(f'/api/export/xlsx?ids={id}')
    book = load_workbook(io.BytesIO(xlsx.content))
    assert book.active.max_row == 4
    recollect = client.post(f'/api/products/{id}/recollect').json()['id']
    assert wait(client, recollect)['counts']['done'] == 1
    assert client.get('/api/products').json()['total'] == 2
    assert client.post('/api/products/delete', json={'ids':[id]}).json()['deleted'] == 1
    assert client.get(f'/api/products/{id}').status_code == 404

def test_pause_resume_and_cancel(client):
    id = start(client, end_page=10)
    assert client.post(f'/api/jobs/{id}/pause').status_code == 200
    time.sleep(.1)
    job = next(j for j in client.get('/api/jobs').json() if j['id'] == id)
    assert job['status'] == 'paused'
    assert client.post(f'/api/jobs/{id}/resume').status_code == 200
    assert client.post(f'/api/jobs/{id}/cancel').status_code == 200
    assert wait(client, id)['status'] == 'cancelled'
    assert client.post(f'/api/jobs/{id}/resume').status_code == 409

def test_failure_retry(client, monkeypatch):
    original = MockProvider.detail
    async def fail(self, seed, config):
        raise ProviderError('테스트 실패')
    monkeypatch.setattr(MockProvider, 'detail', fail)
    id = start(client, page_size=2)
    job = wait(client, id)
    assert job['status'] == 'completed_with_errors' and job['counts']['failed'] == 2
    assert any(l['level'] == 'error' for l in client.get(f'/api/jobs/{id}/logs').json())
    monkeypatch.setattr(MockProvider, 'detail', original)
    assert client.post(f'/api/jobs/{id}/retry').status_code == 200
    assert wait(client, id)['counts']['done'] == 2

def test_crash_recovery_preserves_completed():
    # No live worker is running here; simulate persisted interrupted work directly.
    from app.db import Base, engine
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with Session.begin() as db:
        j = Job(source='mock', keyword='복구', config={'keyword':'복구','start_page':1,'end_page':1,'page_size':2,'currency':'USD','country':'KR'}, next_page=2,status='running')
        db.add(j); db.flush(); id=j.id
        db.add(JobItem(job_id=id, external_id='done',seed={'product_id':'done'},status='done'))
        db.add(JobItem(job_id=id, external_id='pending',seed={'product_id':'pending'},status='running'))
    recover()
    with Session.begin() as db:
        assert db.get(Job,id).status == 'paused'
        statuses = list(db.scalars(select(JobItem.status).where(JobItem.job_id==id).order_by(JobItem.id)))
        assert statuses == ['done','pending']
        db.get(Job,id).status='queued'
    asyncio.run(process_one(id,MockProvider()))
    asyncio.run(process_one(id,MockProvider()))
    with Session() as db:
        assert db.get(Job,id).status == 'completed'
        assert db.scalar(select(func.count()).select_from(Product)) == 1

def test_missing_fields_and_no_cartesian_product():
    product = affiliate_product({'product_id':123,'original_price':'3','original_price_currency':'USD'}, 'KRW')
    assert product['currency'] == 'USD' and product['brand'] is None
    result = {'ae_item_base_info_dto':{},'ae_item_sku_info_dtos':{'ae_item_sku_info_d_t_o':[
        {'sku_id':'one','sku_price':'4','sku_available_stock':0,'ae_sku_property_dtos':{'ae_sku_property_d_t_o':[
            {'sku_property_name':'Color','sku_property_value':'Black'},{'sku_property_name':'Size','sku_property_value':'M3'}]}},
        {'sku_id':'two','ae_sku_property_dtos':{'ae_sku_property_d_t_o':[{'sku_property_name':'Color','sku_property_value':'White'}]}}
    ]}}
    enrich_ds(product,result)
    assert product['option_status'] == 'partial'
    assert len(product['variants']) == 2
    assert product['variants'][0]['options'][0]['value'] == 'Black'
    assert product['variants'][0]['stock'] == 0 and product['variants'][0]['available'] is None
    assert product['variants'][1]['stock'] is None

def test_retry_after_seconds_and_http_date():
    assert retry_delay('120',0) == 120
    date = format_datetime(datetime.now(timezone.utc)+timedelta(seconds=120))
    assert 118 <= retry_delay(date,0) <= 120
    assert 4 <= retry_delay(None,2) < 5

def test_api_429_and_signing(monkeypatch):
    monkeypatch.setenv('ALI_APP_KEY','test-key')
    monkeypatch.setenv('ALI_APP_SECRET','test-secret')
    monkeypatch.setenv('API_INTERVAL_SECONDS','0')
    monkeypatch.setenv('API_MAX_RETRIES','1')
    calls=[]
    async def fake_post(self,url,data):
        calls.append(data)
        if len(calls)==1:
            return httpx.Response(429,headers={'Retry-After':'0'})
        return httpx.Response(200,json={'aliexpress_affiliate_product_query_response':{'resp_result':{'resp_code':200,'result':{'products':{'product':[]}}}}})
    monkeypatch.setattr(httpx.AsyncClient,'post',fake_post)
    rows=asyncio.run(OfficialProvider().search({'keyword':'나사','page_size':50,'country':'KR','currency':'KRW','sort':''},1))
    assert rows == [] and len(calls)==2 and calls[0]['sign']
    assert calls[0]['ship_to_country']=='KR'

def test_input_guards_and_live_credentials(client,monkeypatch):
    monkeypatch.delenv('ALI_APP_KEY',raising=False)
    monkeypatch.delenv('ALI_APP_SECRET',raising=False)
    assert client.post('/api/jobs',json={'keyword':'나사','source':'live'}).status_code==400
    assert client.post('/api/jobs',json={'keyword':' ','source':'mock'}).status_code==422
    assert client.post('/api/jobs',json={'keyword':'나사','start_page':10,'end_page':1}).status_code==422
    assert client.post('/api/jobs',json={'keyword':'나사','page_size':51}).status_code==422
    assert client.post('/api/jobs',json={'keyword':'나사'},headers={'Origin':'https://untrusted.example'}).status_code==403
    assert client.get('/api/products?page=0').status_code==422

def test_csv_formula_protection():
    from app.main import safe_cell
    assert safe_cell(' =HYPERLINK("bad")').startswith("'")
    assert safe_cell(None)==''

def test_image_rejects_untrusted_url():
    from app.images import save_image
    with pytest.raises(ValueError):
        asyncio.run(save_image('live','x',{'url':'http://127.0.0.1/private'}))
